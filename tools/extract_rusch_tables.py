#!/usr/bin/env python3
"""Extrai as tabelas de Rüsch do workbook do IME para um JSON versionado.

A planilha de origem (`Tabelas Rusch Final - IME.xls`, formato BIFF legado) não é
versionada. Este script consome a versão `.xlsx` já convertida e produz um único
artefato determinístico:

  crates/freerusch-core/data/rusch_tables.json

O crate `freerusch-core` embute esse JSON com `include_str!` e o desserializa uma
vez (`rusch_tables()`). Gerar código Rust com 56 mil coeficientes foi descartado:
um arquivo de ~1,6 MB de expressões degrada o tempo de compilação sem ganho.

Regras de fidelidade:
  * nenhum valor é interpolado, arredondado ou reinterpretado;
  * célula `-` vira ausência (`null`), nunca `0,0`;
  * célula com erro de planilha (`=#REF!`) ou número ilegível (`2,.82`) vira
    ausência **e** entra na lista `warnings` do JSON, para conferência manual;
  * rótulos de coluna são preservados crus (inclusive prefixos `¾`, `(+)`, `-`, `_`),
    porque o significado desses prefixos ainda não foi confirmado com o documento
    original.

Conversão do `.xls` (uma vez, com o LibreOffice empacotado):
  node <libreoffice-kit>/cli.js convert \
      --input "Tabelas Rusch Final - IME.xls" --output target/ref/tabelas-rusch.xlsx

Uso:
  python tools/extract_rusch_tables.py --xlsx target/ref/tabelas-rusch.xlsx --all
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

# Amostra usada enquanto o pipeline era validado. O build real roda com `--all`,
# que descobre todas as planilhas de tabela do workbook.
SAMPLE_SHEETS = ["Nr 1.1", "Nr 4.1", "Nr 16.1", "Nr 99", "Nr 27.1"]

SHEET_PATTERN = re.compile(r"^Nr\.?\s*\d+")
SHEET_NAME = re.compile(r"^Nr\.?\s*(\d+)(?:\.(\d+))?$")

DATA_ROWS = 13  # lx/a = 0.5, 1, 1.5, 2, 2.5, 3, 4, 5, 6, 7, 8, 9, 10

REPO = Path(__file__).resolve().parent.parent
OUT_JSON = REPO / "crates/freerusch-core/data/rusch_tables.json"


DASHES = {"-", "–", "—", ""}

# Células que o próprio workbook carrega como erro de planilha (fórmula
# quebrada). Não são valor nem "não se aplica": viram ausência **registrada**,
# para que a lacuna apareça na auditoria em vez de virar um número inventado.
EXCEL_ERRORS = ("#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A")


def as_optional_float(value, where: str, warnings: list):
    """Converte uma célula em float, ou `None` para ausência.

    Nas tabelas, `-` significa que o coeficiente não se aplica (contribuição
    nula). Isso é preservado como ausência explícita, nunca como 0,0 inventado.
    Erros de planilha (`=#REF!` etc.) também viram ausência, mas ficam
    registrados em `warnings` e no JSON para revisão.
    """
    if isinstance(value, bool):
        raise SystemExit(f"célula booleana em {where}")
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if text in DASHES:
        return None
    if any(err in text for err in EXCEL_ERRORS):
        warnings.append({"cell": where, "raw": text, "reason": "spreadsheet error"})
        return None
    if text.startswith("="):
        raise SystemExit(f"fórmula sem valor em cache em {where}: {value!r}")
    try:
        return float(text.replace(",", "."))
    except ValueError:
        # Erro de digitação na planilha (ex.: "2,.82"). Reparar o número seria
        # inventar um coeficiente: fica ausente e registrado para conferência
        # contra o documento original.
        warnings.append({"cell": where, "raw": text, "reason": "unparseable number"})
        return None


def find_header(grid):
    for r, row in enumerate(grid):
        for c, v in enumerate(row):
            if isinstance(v, str) and v.strip().lower() == "lx/a":
                return r, c
    return None, None


def parse_id(name: str) -> tuple[int | None, int | None]:
    """Número e variante derivados do NOME da planilha (autoritativo)."""
    match = SHEET_NAME.match(name.strip())
    if not match:
        return None, None
    return int(match.group(1)), (
        int(match.group(2)) if match.group(2) else None
    )


def build_grid(ws_formula, ws_values, warnings: list) -> list:
    """Grade resolvida: fórmulas substituídas pelo valor em cache.

    Algumas planilhas trazem fórmulas reais (`=C12*1.97`). Ler apenas em modo
    `data_only=False` devolveria o texto da fórmula, não o número. Cruzamos as
    duas leituras: a estrutura vem da planilha de fórmulas, os valores das
    células calculadas vêm da leitura `data_only=True`.
    """
    raw = [[c.value for c in row] for row in ws_formula.iter_rows(max_row=40, max_col=40)]
    cached = [[c.value for c in row] for row in ws_values.iter_rows(max_row=40, max_col=40)]

    resolved = []
    for r, row in enumerate(raw):
        out = []
        for c, cell in enumerate(row):
            if isinstance(cell, str) and cell.startswith("="):
                value = None
                if r < len(cached) and c < len(cached[r]):
                    value = cached[r][c]
                if value is None:
                    warnings.append(
                        {
                            "cell": f"r{r + 1}c{c + 1}",
                            "raw": cell,
                            "reason": "formula without cached value",
                        }
                    )
                out.append(value)
            else:
                out.append(cell)
        resolved.append(out)
    return resolved


def parse_sheet(name, grid, warnings: list | None = None):
    warnings = warnings if warnings is not None else []
    header_row, label_col = find_header(grid)
    if header_row is None:
        raise SystemExit(f"{name!r}: linha 'lx/a' não encontrada")
    if header_row + 3 >= len(grid):
        raise SystemExit(f"{name!r}: cabeçalho incompleto")

    group_row = grid[header_row]
    var_row = grid[header_row + 1]
    ta_row = grid[header_row + 2]
    letter_row = grid[header_row + 3]

    columns = []
    group = ""
    for c in range(label_col + 1, max(len(r) for r in grid[: header_row + 5])):
        g = group_row[c] if c < len(group_row) else None
        if isinstance(g, str) and g.strip():
            group = g.strip()
        var = var_row[c] if c < len(var_row) else None
        ta = ta_row[c] if c < len(ta_row) else None
        letter = letter_row[c] if c < len(letter_row) else None
        # A linha das letras (L / p / p') é o discriminador: colunas `p'` não
        # têm `t/a` próprio (herdam "für alle Werte t/a"), mas têm letra.
        if not (isinstance(letter, str) and letter.strip()):
            continue
        kind = "wheel" if isinstance(letter, str) and letter.strip() == "L" else "uniform"
        columns.append(
            {
                "index": c,
                "group": group,
                "kind": kind,
                "letter": (letter.strip() if isinstance(letter, str) else ""),
                "ta": (float(ta) if isinstance(ta, (int, float)) and not isinstance(ta, bool) else None),
                "values": [],
            }
        )

    if not columns:
        raise SystemExit(f"{name!r}: nenhuma coluna de coeficientes reconhecida")

    rows = []
    for r in range(header_row + 4, len(grid)):
        row = grid[r]
        label = row[label_col] if label_col < len(row) else None
        if not isinstance(label, (int, float)) or isinstance(label, bool):
            if rows:
                break
            continue
        rows.append(float(label))
        for col in columns:
            c = col["index"]
            cell = row[c] if c < len(row) else None
            col["values"].append(
                as_optional_float(cell, f"{name}!r{r + 1}c{c + 1}", warnings)
            )

    if len(rows) != DATA_ROWS:
        raise SystemExit(f"{name!r}: {len(rows)} linhas de dados, esperado {DATA_ROWS}")
    for col in columns:
        if len(col["values"]) != DATA_ROWS:
            raise SystemExit(f"{name!r}: coluna {col['group']!r} incompleta")

    # Identidade: o nome da planilha é autoritativo. O rótulo interno é só
    # conferido — há planilhas cujo cabeçalho repete o número da anterior.
    number, variant = parse_id(name)
    label_raw = None
    for row in grid[:header_row]:
        for v in row:
            if isinstance(v, str) and SHEET_NAME.match(v.strip()):
                label_raw = v.strip()
    if label_raw and parse_id(label_raw) != (number, variant):
        warnings.append(
            {
                "cell": "header",
                "raw": label_raw,
                "reason": f"sheet name {name!r} disagrees with header label",
            }
        )

    param = ratio_raw = None
    for row in grid[:header_row]:
        for c, v in enumerate(row):
            if isinstance(v, str) and re.match(r"^(ly/lx|l/a)\s*=", v.strip()):
                param = v.strip().split("=")[0].strip()
                for k in range(c + 1, len(row)):
                    if row[k] is not None:
                        ratio_raw = str(row[k]).strip()
                        break

    ratio = None if ratio_raw in (None, "∞") else float(ratio_raw.replace(",", "."))

    return {
        "sheet": name,
        "number": number,
        "variant": variant,
        "param": param,
        "ratio_raw": ratio_raw,
        "ratio": ratio,
        "ratio_is_infinite": ratio_raw == "∞",
        "rows": rows,
        "columns": columns,
        "warnings": warnings,
    }


def sort_key(name: str) -> tuple[int, int, str]:
    """Ordem determinística: número, variante, nome."""
    match = SHEET_NAME.match(name.strip())
    if not match:
        return (10**6, 10**6, name)
    return (int(match.group(1)), int(match.group(2) or 0), name)


def discover_sheets(workbook) -> list[str]:
    """Todas as planilhas de tabela do workbook, em ordem determinística."""
    found = [name for name in workbook.sheetnames if SHEET_PATTERN.match(name.strip())]
    return sorted(found, key=sort_key)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", required=True, help="workbook .xlsx convertido")
    parser.add_argument(
        "--all",
        action="store_true",
        help="extrai todas as planilhas Nr* (padrão: apenas a amostra)",
    )
    parser.add_argument("--source-name", default="Tabelas Rusch Final - IME.xls")
    parser.add_argument("--source-sha256", default="")
    args = parser.parse_args()

    xlsx = Path(args.xlsx)
    workbook = openpyxl.load_workbook(xlsx, read_only=True, data_only=False)
    workbook_values = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    sheets = discover_sheets(workbook) if args.all else SAMPLE_SHEETS

    missing = [s for s in sheets if s not in workbook.sheetnames]
    if missing:
        raise SystemExit(f"planilhas ausentes: {missing}")

    tables = []
    for name in sheets:
        sheet_warnings: list = []
        grid = build_grid(workbook[name], workbook_values[name], sheet_warnings)
        table = parse_sheet(name, grid, sheet_warnings)
        tables.append(table)
        kinds = {}
        for col in table["columns"]:
            kinds[col["letter"]] = kinds.get(col["letter"], 0) + 1
        dashes = sum(1 for c in table["columns"] for v in c["values"] if v is None)
        print(
            f"  {name:<10} nr={table['number']}.{table['variant']} "
            f"{table['param']}={table['ratio_raw']} "
            f"cols={len(table['columns'])} {kinds} sem_valor={dashes}"
        )
    workbook.close()
    workbook_values.close()

    # Identidade única: (number, variant) é a chave de busca pública.
    seen = {}
    for table in tables:
        key = (table["number"], table["variant"])
        if key in seen:
            raise SystemExit(
                f"identidade duplicada {key}: {seen[key]!r} e {table['sheet']!r}"
            )
        seen[key] = table["sheet"]

    payload = {
        "source": {
            "name": args.source_name,
            "sha256": args.source_sha256,
            "xlsx_sha256": hashlib.sha256(xlsx.read_bytes()).hexdigest().upper(),
            "tool": "tools/extract_rusch_tables.py",
        },
        "rows_axis": {"name": "lx/a", "values": tables[0]["rows"]},
        "row_count": DATA_ROWS,
        "scope": {
            "note": "conjunto completo" if args.all else "recorte parcial",
            "sheet_count": len(sheets),
            "sheets": sheets,
        },
        "tables": tables,
        "warnings": [
            {"sheet": t["sheet"], **w} for t in tables for w in t["warnings"]
        ],
    }

    OUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUT_JSON.write_text(
        json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    total = sum(len(c["values"]) for t in tables for c in t["columns"])
    warn_count = sum(len(t["warnings"]) for t in tables)
    print(f"\ntabelas={len(tables)} valores={total} celulas_com_erro={warn_count}")
    if warn_count:
        print("  células de planilha com erro (viram ausência, registradas no JSON):")
        for t in tables:
            for w in t["warnings"][:3]:
                print(f"    {t['sheet']} {w['cell']} {w['raw']}")
    print(f"escrito: {OUT_JSON.relative_to(REPO)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
